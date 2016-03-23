#!/usr/bin/env python
# -*- coding: utf-8 -*-

import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Inches
import os
import shutil


def span(rowSpan,colSpan,cell_row,cell_col,newTable):
	for i in range(int(rowSpan)):
		for j in range(1,int(colSpan)):
			newTable.cell(cell_row,cell_col).merge(newTable.cell(cell_row+i,cell_col+j))

		
	for i in range(1,int(rowSpan)):
		newTable.cell(cell_row,cell_col).merge(newTable.cell(cell_row+i,cell_col))	
		
if not os.path.isdir("target"):
    os.mkdir("target")

path = os.getcwd() 
for filename in os.listdir(path):
    	if not filename.endswith('.xml'): continue
	document = Document()

        print filename

	tree = ET.parse(filename)
	root = tree.getroot()
	id=0
	for table in root.findall('Table'):

	    rows = table.find('Rows').text
	    rows = int(rows)
	    columns = table.find('Cols').text
	    columns = int(columns)
            if rows == 0 or columns == 0:
                id=id+1
                continue
            #print "Table Id: "+str(id)
            #print "Table Rows,Columns: "+str(rows)+","+str(columns)
	    newTable = document.add_table(rows=rows, cols=columns)	#create table with given rows and columns
	    bid=0
	    for blocks in table.findall('Block'):
		    row = int(root[id][bid+4][0][0].text)
		    col = int(root[id][bid+4][0][1].text)
		    data = blocks.find('Text').text
		    rowSpan = blocks.find('RowSpan').text

		    colSpan = blocks.find('ColSpan').text
		    if rowSpan>1 or colSpan>1:
			span(rowSpan,colSpan,row,col,newTable)
		    if data is not None:
		    	newTable.cell(row,col).text=data.strip()


		    #print "Block id: " + str(bid)
		    bid=bid+1

	    document.add_page_break()

	    id=id+1

        file='target/'+str(filename)

        shutil.move(str(filename),file)
	filename=filename.split('.')

        filename[len(filename)-1] = "docx"

        filename = ".".join([ str(item) for item in filename])

        document.save('target/'+filename)




